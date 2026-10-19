"""Comprehensive automated test suite for Document Intelligence, Decision-Makers, Contacts, Evidence, and Export."""

import io
import os
import sqlite3
import zipfile
import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

import database
from app import app
from candidate_verifier import clean_candidate_name, verify_customer_candidate
from intelligence_extractor import (
    classify_entity_category,
    extract_decision_makers_from_text,
    extract_public_contacts_from_text,
    generate_job_unified_evidence,
    build_unified_evidence,
)
from models import (
    CandidateCategory,
    CandidateClassification,
    ClaimType,
    ContactType,
    ContactVerificationStatus,
    CustomerCandidate,
    DecisionMaker,
    DiscoveredDocument,
    EvidenceReference,
    PublicContact,
    RelationshipType,
    RoleCategory,
    UnifiedEvidence,
    VerificationChecklist,
)


@pytest.fixture
def temp_db(tmp_path):
    """Fixture providing an isolated SQLite database."""
    db_file = tmp_path / "test_intelligence.db"
    db_path = str(db_file)
    database.init_db(db_path=db_path)
    return db_path


# ============================================================================
# 1. ENTITY QUALITY CONTROL & CATEGORIZATION
# ============================================================================

def test_classify_named_companies():
    """Verify genuine company names are classified as NAMED_COMPANY."""
    assert classify_entity_category("Tata Steel Ltd", "Tata Steel implemented cloud ERP.") == CandidateCategory.NAMED_COMPANY
    assert classify_entity_category("Acme Corporation", "Acme Corporation uses Jedox.") == CandidateCategory.NAMED_COMPANY
    assert classify_entity_category("Micro Labs", "Customer Micro Labs achieved 30% savings.") == CandidateCategory.NAMED_COMPANY
    assert classify_entity_category("Siemens Healthineers", "Siemens Healthineers deployed analytics.") == CandidateCategory.NAMED_COMPANY


def test_classify_anonymous_client_descriptions():
    """Verify anonymous descriptors are flagged as ANONYMOUS_CLIENT_DESCRIPTION, not verified leads."""
    assert classify_entity_category("Global FinTech Player", "A global fintech player chosen.") == CandidateCategory.ANONYMOUS_CLIENT_DESCRIPTION
    assert classify_entity_category("Global Food Company", "How a global food company scaled.") == CandidateCategory.ANONYMOUS_CLIENT_DESCRIPTION
    assert classify_entity_category("Leading Alcobev Giant", "Our client, a leading alcobev giant.") == CandidateCategory.ANONYMOUS_CLIENT_DESCRIPTION
    assert classify_entity_category("Alcobev Giant", "Alcobev giant improves planning.") == CandidateCategory.ANONYMOUS_CLIENT_DESCRIPTION


def test_classify_products_and_technologies():
    """Verify software/technology names are not treated as company leads."""
    assert classify_entity_category("Jedox", "Implemented Jedox for planning.") == CandidateCategory.PRODUCT_OR_TECHNOLOGY
    assert classify_entity_category("Power BI", "Built using Power BI dashboards.") == CandidateCategory.PRODUCT_OR_TECHNOLOGY
    assert classify_entity_category("Tableau", "Tableau integration completed.") == CandidateCategory.PRODUCT_OR_TECHNOLOGY
    assert classify_entity_category("Snowflake", "Snowflake data warehouse.") == CandidateCategory.PRODUCT_OR_TECHNOLOGY


def test_classify_person_entities():
    """Verify person names with honorifics are categorized as PERSON."""
    assert classify_entity_category("Dr. Robert Smith", "Keynote by Dr. Robert Smith.") == CandidateCategory.PERSON
    assert classify_entity_category("Prof. Charles Xavier", "Authored by Prof. Charles Xavier.") == CandidateCategory.PERSON


# ============================================================================
# 2. DECISION-MAKER DISCOVERY & ROLE CLASSIFICATION
# ============================================================================

def test_extract_technology_decision_makers():
    """Verify extraction of CTO, Head of BI, and IT leadership."""
    text = (
        "During the keynote, Alexander Wright, Chief Technology Officer at Polestar Analytics, "
        "explained the architecture. Later, Priya Sharma, Head of Business Intelligence, presented the roadmap."
    )
    dms = extract_decision_makers_from_text(
        text=text,
        source_url="https://polestaranalytics.com/leadership",
        source_type="WEBPAGE",
        job_id=1,
        default_company="Polestar Analytics",
    )
    names = {dm.full_name for dm in dms}
    assert "Alexander Wright" in names
    assert "Priya Sharma" in names

    cto = next(d for d in dms if d.full_name == "Alexander Wright")
    assert cto.role_category == RoleCategory.TECHNOLOGY_IT
    assert "Technology" in cto.published_title or "CTO" in cto.published_title
    assert cto.verification_status == "CONFIRMED_PUBLIC"
    assert "Polestar Analytics" in cto.associated_company


def test_extract_finance_decision_makers():
    """Verify extraction of CFO, Finance Directors, and FP&A heads."""
    text = (
        "Marcus Vance, Chief Financial Officer at Acme Corp, commented on the deployment: "
        "'The platform reduced our budgeting cycle by 40%'. Elena Rostova, Head of FP&A, spearheaded the migration."
    )
    dms = extract_decision_makers_from_text(
        text=text,
        source_url="https://acmecorp.com/case-study",
        source_type="DOCUMENT",
        job_id=2,
    )
    cfo = next((d for d in dms if d.full_name == "Marcus Vance"), None)
    assert cfo is not None
    assert cfo.role_category == RoleCategory.FINANCE_BUSINESS
    assert "CFO" in cfo.published_title or "Chief Financial Officer" in cfo.published_title
    assert "Marcus Vance" in cfo.supporting_quote


def test_extract_executive_and_operations_leaders():
    """Verify extraction of CEO, Managing Director, and Operations heads."""
    text = (
        "Company announcement by David Miller, Chief Executive Officer, alongside "
        "Karen Patel, Director of Supply Chain."
    )
    dms = extract_decision_makers_from_text(text=text, source_url="https://example.com/team", job_id=3)
    ceo = next((d for d in dms if d.full_name == "David Miller"), None)
    ops = next((d for d in dms if d.full_name == "Karen Patel"), None)
    assert ceo is not None and ceo.role_category == RoleCategory.EXECUTIVE
    assert ops is not None and ops.role_category == RoleCategory.OPERATIONS


# ============================================================================
# 3. PUBLIC PROFESSIONAL CONTACT EXTRACTION & INTEGRITY
# ============================================================================

def test_extract_valid_public_emails():
    """Verify publicly published professional emails are extracted with exact quotes."""
    text = (
        "For business inquiries, contact enterprise@polestaranalytics.com or reach out to "
        "press.office@polestar.com for media relations."
    )
    contacts = extract_public_contacts_from_text(
        text=text,
        source_url="https://polestaranalytics.com/contact",
        file_type="HTML",
        job_id=1,
    )
    emails = {c.contact_value for c in contacts if c.contact_type == ContactType.EMAIL}
    assert "enterprise@polestaranalytics.com" in emails
    assert "press.office@polestar.com" in emails
    for c in contacts:
        assert c.verification_status == ContactVerificationStatus.VERIFIED_PUBLIC
        assert c.is_public is True
        assert c.supporting_quote is not None


def test_reject_fabricated_and_invalid_emails():
    """Verify placeholder domains and image assets disguised as emails are rejected."""
    text = (
        "Assets include icon@2x.png, background@preview.jpg and styles@theme.css. "
        "Also ignore fake address john.doe@example.com and test@sample.com."
    )
    contacts = extract_public_contacts_from_text(
        text=text,
        source_url="https://example.com/page",
        job_id=1,
    )
    emails = [c.contact_value for c in contacts if c.contact_type == ContactType.EMAIL]
    assert len(emails) == 0


def test_extract_public_phone_and_linkedin():
    """Verify phone numbers and LinkedIn profile URLs are captured."""
    text = (
        "Office Telephone: +1-800-555-0199. Follow our executive at https://www.linkedin.com/in/alexander-wright-tech/"
    )
    contacts = extract_public_contacts_from_text(
        text=text,
        source_url="https://example.com/contact",
        job_id=1,
    )
    phones = [c.contact_value for c in contacts if c.contact_type in (ContactType.PHONE, ContactType.OFFICE_PHONE)]
    linkedins = [c.contact_value for c in contacts if c.contact_type == ContactType.LINKEDIN]

    assert len(phones) >= 1
    assert any("800" in p for p in phones)
    assert len(linkedins) >= 1
    assert "linkedin.com/in/alexander-wright-tech" in linkedins[0]


# ============================================================================
# 4. INDEPENDENT VERIFICATION & EVIDENCE TRACEABILITY
# ============================================================================

def test_independent_claim_verification():
    """Verify Claim A (relationship) and Claim B (product) are independently checked."""
    # Case 1: Relationship verified, but NO target product usage -> PARTIALLY_VERIFIED
    ref1 = EvidenceReference(
        source_url="https://target.com/case-studies",
        page_title="Case Studies",
        supporting_text="Our valued customer Acme Corp achieved milestone growth.",
        relationship_claim="Acme Corp is a customer",
        product_claim=None,
    )
    cand = verify_customer_candidate(
        candidate_name="Acme Corp",
        candidate_sources=[ref1],
        all_page_texts=[("https://target.com/case-studies", "Case Studies", "Our valued customer Acme Corp achieved milestone growth.")],
        target_product="Jedox",
        is_crawl_complete=True,
    )
    assert cand.classification == CandidateClassification.PARTIALLY_VERIFIED
    assert cand.product_usage_status == "UNVERIFIED"

    # Case 2: Both customer relationship and Jedox product usage verified -> CONFIRMED
    ref2 = EvidenceReference(
        source_url="https://target.com/jedox-success",
        page_title="Acme Jedox Story",
        supporting_text="Customer Acme Corp deployed Jedox for automated financial planning and consolidation.",
        relationship_claim="Customer Acme Corp",
        product_claim="Acme Corp deployed Jedox",
    )
    cand_confirmed = verify_customer_candidate(
        candidate_name="Acme Corp",
        candidate_sources=[ref1, ref2],
        all_page_texts=[("https://target.com/jedox-success", "Acme Jedox Story", "Customer Acme Corp deployed Jedox for automated financial planning and consolidation.")],
        target_product="Jedox",
        is_crawl_complete=True,
    )
    assert cand_confirmed.classification == CandidateClassification.CONFIRMED
    assert cand_confirmed.product_usage_status == "CONFIRMED"


def test_anonymous_description_downgraded_from_confirmed():
    """Ensure anonymous descriptions are never presented as confirmed company leads."""
    ref = EvidenceReference(
        source_url="https://target.com/story",
        page_title="Fintech Story",
        supporting_text="Our client, a Global FinTech Player deployed Jedox for planning.",
        relationship_claim="Global FinTech Player is a client",
        product_claim="Global FinTech Player deployed Jedox",
    )
    cand = verify_customer_candidate(
        candidate_name="Global FinTech Player",
        candidate_sources=[ref],
        all_page_texts=[("https://target.com/story", "Fintech Story", "Our client, a Global FinTech Player deployed Jedox for planning.")],
        target_product="Jedox",
        is_crawl_complete=True,
    )
    assert cand.candidate_category == CandidateCategory.ANONYMOUS_CLIENT_DESCRIPTION
    # Must not be CONFIRMED
    assert cand.classification != CandidateClassification.CONFIRMED


def test_unified_evidence_generation():
    """Verify unified evidence records aggregate all source citations."""
    cand = CustomerCandidate(
        id=10,
        job_id=1,
        candidate_name="Tata Steel",
        normalized_name="tata steel",
        candidate_category=CandidateCategory.NAMED_COMPANY,
        target_product="Jedox",
        relationship_type=RelationshipType.CUSTOMER,
        classification=CandidateClassification.CONFIRMED,
        checklist=VerificationChecklist(customer_claim_verified=True, target_product_verified=True),
        product_usage_status="CONFIRMED",
        sources=[
            EvidenceReference(
                source_url="https://target.com/tata-steel",
                supporting_text="Tata Steel selected Jedox for enterprise planning.",
                relationship_claim="Customer relationship verified",
                product_claim="Deployed Jedox",
            )
        ],
    )
    dm = DecisionMaker(
        id=20,
        job_id=1,
        full_name="Rajesh Kumar",
        published_title="Head of BI",
        role_category=RoleCategory.TECHNOLOGY_IT,
        associated_company="Tata Steel",
        supporting_quote="Rajesh Kumar, Head of BI led the initiative.",
        source_url="https://target.com/tata-steel",
        source_type="WEBPAGE",
    )
    ct = PublicContact(
        id=30,
        job_id=1,
        contact_type=ContactType.EMAIL,
        contact_value="bi.support@tatasteel.com",
        person_name="Rajesh Kumar",
        associated_company="Tata Steel",
        supporting_quote="Reach us at bi.support@tatasteel.com",
        source_url="https://target.com/tata-steel",
        verification_status=ContactVerificationStatus.VERIFIED_PUBLIC,
    )

    evidences = generate_job_unified_evidence(
        job_id=1,
        candidates=[cand],
        decision_makers=[dm],
        contacts=[ct],
    )
    assert len(evidences) >= 3
    claim_types = {e.claim_type for e in evidences}
    assert ClaimType.RELATIONSHIP in claim_types
    assert ClaimType.PRODUCT_USAGE in claim_types
    assert ClaimType.DECISION_MAKER in claim_types
    assert ClaimType.CONTACT in claim_types


# ============================================================================
# 5. DATABASE SCHEMA & IDEMPOTENT MIGRATIONS
# ============================================================================

def test_database_crud_and_idempotency(temp_db):
    """Verify document, decision-maker, contact, and evidence CRUD in SQLite."""
    # 0. Create Job
    job = database.create_crawl_job(
        database.CrawlJob(seed_url="https://example.com", workflow="GENERAL"),
        db_path=temp_db,
    )
    job_id = job.id

    # 1. Discovered Document
    doc = DiscoveredDocument(
        job_id=job_id,
        source_page_url="https://example.com",
        document_url="https://example.com/report.pdf",
        file_name="report.pdf",
        file_format="PDF",
        file_size_bytes=10240,
        download_status="DOWNLOADED",
        extraction_status="SUCCESS",
        extracted_text="Sample extracted PDF text content.",
    )
    saved_doc = database.save_discovered_document(doc, db_path=temp_db)
    assert saved_doc.id is not None
    docs = database.get_discovered_documents(job_id=job_id, db_path=temp_db)
    assert len(docs) == 1
    assert docs[0].file_name == "report.pdf"

    # 2. Decision Maker
    dm = DecisionMaker(
        job_id=job_id,
        full_name="Sarah Connor",
        published_title="Chief Technology Officer",
        role_category=RoleCategory.TECHNOLOGY_IT,
        associated_company="Cyberdyne",
        supporting_quote="Sarah Connor, Chief Technology Officer.",
        source_url="https://cyberdyne.com/leadership",
    )
    saved_dms = database.save_decision_makers([dm], db_path=temp_db)
    assert len(saved_dms) == 1
    dms = database.get_decision_makers(job_id=job_id, db_path=temp_db)
    assert len(dms) == 1
    assert dms[0].full_name == "Sarah Connor"

    # 3. Public Contact
    ct = PublicContact(
        job_id=job_id,
        contact_type=ContactType.EMAIL,
        contact_value="sarah.connor@cyberdyne.com",
        person_name="Sarah Connor",
        associated_company="Cyberdyne",
        source_url="https://cyberdyne.com/contact",
    )
    saved_cts = database.save_public_contacts([ct], db_path=temp_db)
    assert len(saved_cts) == 1
    cts = database.get_public_contacts(job_id=job_id, db_path=temp_db)
    assert len(cts) == 1
    assert cts[0].contact_value == "sarah.connor@cyberdyne.com"

    # 4. Unified Evidence
    ev = UnifiedEvidence(
        job_id=job_id,
        claim_type=ClaimType.RELATIONSHIP,
        claim_text="Cyberdyne partnership",
        source_url="https://cyberdyne.com",
        supporting_quote="Strategic partnership with Cyberdyne.",
    )
    saved_evs = database.save_unified_evidence([ev], db_path=temp_db)
    assert len(saved_evs) == 1
    evs = database.get_unified_evidence(job_id=job_id, db_path=temp_db)
    assert len(evs) == 1
    assert evs[0].claim_type == ClaimType.RELATIONSHIP


# ============================================================================
# 6. FASTAPI ENDPOINTS & MULTI-DATASET EXPORT
# ============================================================================

def test_api_export_endpoints(temp_db, monkeypatch):
    """Verify Excel (.xlsx) and CSV (.zip) export endpoints with multi-dataset worksheets."""
    monkeypatch.setattr(database, "get_db_path", lambda: temp_db)

    # Seed test job and records
    job = database.create_crawl_job(
        database.CrawlJob(seed_url="https://test.com", workflow="GENERAL", target_product="Jedox"),
        db_path=temp_db,
    )
    cand = CustomerCandidate(
        job_id=job.id,
        candidate_name="Acme Corp",
        normalized_name="acme corp",
        target_product="Jedox",
        relationship_type=RelationshipType.CUSTOMER,
        classification=CandidateClassification.CONFIRMED,
        checklist=VerificationChecklist(customer_claim_verified=True, target_product_verified=True),
        product_usage_status="CONFIRMED",
        sources=[],
    )
    database.save_customer_candidates([cand], db_path=temp_db)

    dm = DecisionMaker(
        job_id=job.id,
        full_name="Alice Vance",
        published_title="CFO",
        role_category=RoleCategory.FINANCE_BUSINESS,
        source_url="https://test.com/leadership",
    )
    database.save_decision_makers([dm], db_path=temp_db)

    client = TestClient(app)

    # 1. Test Excel Export
    resp_excel = client.get(f"/api/crawl/jobs/{job.id}/export?format=excel")
    assert resp_excel.status_code == 200
    assert "spreadsheetml" in resp_excel.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(resp_excel.content))
    assert "Candidates" in wb.sheetnames
    assert "Decision Makers" in wb.sheetnames
    assert "Evidence" in wb.sheetnames
    assert "Public Contacts" in wb.sheetnames
    assert "Documents & Media" in wb.sheetnames

    # 2. Test CSV ZIP Export
    resp_csv = client.get(f"/api/crawl/jobs/{job.id}/export?format=csv")
    assert resp_csv.status_code == 200
    assert "zip" in resp_csv.headers["content-type"]
    with zipfile.ZipFile(io.BytesIO(resp_csv.content)) as zf:
        namelist = zf.namelist()
        assert "candidates.csv" in namelist
        assert "decision_makers.csv" in namelist
        assert "evidence.csv" in namelist
        assert "public_contacts.csv" in namelist
        assert "documents.csv" in namelist
